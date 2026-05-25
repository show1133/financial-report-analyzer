# -*- coding: utf-8 -*-
"""
Walmart 跨時代財報分析報告生成器
結合孫子兵法「道天地將法」五構面框架
"""

from docx import Document
from docx.shared import Inches, Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os

def set_cell_shading(cell, color_hex):
    shading = OxmlElement('w:shd')
    shading.set(qn('w:fill'), color_hex)
    shading.set(qn('w:val'), 'clear')
    cell._tc.get_or_add_tcPr().append(shading)

def set_table_border(table):
    tbl = table._tbl
    tblPr = tbl.tblPr if tbl.tblPr is not None else OxmlElement('w:tblPr')
    borders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'single')
        border.set(qn('w:sz'), '4')
        border.set(qn('w:space'), '0')
        border.set(qn('w:color'), '666666')
        borders.append(border)
    tblPr.append(borders)

def add_styled_table(doc, headers, rows, col_widths=None, header_color='1F4E79'):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_border(table)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        for paragraph in cell.paragraphs:
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.name = 'Microsoft JhengHei'
                run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
        set_cell_shading(cell, header_color)
    for row_idx, row_data in enumerate(rows):
        for col_idx, cell_text in enumerate(row_data):
            cell = table.rows[row_idx + 1].cells[col_idx]
            cell.text = str(cell_text)
            for paragraph in cell.paragraphs:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if col_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in paragraph.runs:
                    run.font.size = Pt(9.5)
                    run.font.name = 'Microsoft JhengHei'
                    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
            if row_idx % 2 == 1:
                set_cell_shading(cell, 'F2F2F2')
    if col_widths:
        for i, width in enumerate(col_widths):
            for row in table.rows:
                row.cells[i].width = Cm(width)
    return table

def set_run_font(run, size=11, bold=False, color=None, font_name='Microsoft JhengHei'):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), font_name)
    if color:
        run.font.color.rgb = RGBColor(*color)

def add_heading_styled(doc, text, level=1):
    heading = doc.add_heading(level=level)
    run = heading.add_run(text)
    run.font.name = 'Microsoft JhengHei'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Microsoft JhengHei')
    if level == 1:
        run.font.size = Pt(18)
        run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
    elif level == 2:
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
    elif level == 3:
        run.font.size = Pt(12)
        run.font.color.rgb = RGBColor(0x37, 0x84, 0x5D)
    return heading

def add_body_text(doc, text, indent=False):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.left_indent = Cm(1.0 + level * 0.5)
    p.clear()
    run = p.add_run(text)
    set_run_font(run, size=10.5)
    return p

def add_quote_box(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.5)
    p.paragraph_format.right_indent = Cm(1.5)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.4
    pPr = p._element.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    left_border = OxmlElement('w:left')
    left_border.set(qn('w:val'), 'single')
    left_border.set(qn('w:sz'), '18')
    left_border.set(qn('w:space'), '8')
    left_border.set(qn('w:color'), '2E75B6')
    pBdr.append(left_border)
    pPr.append(pBdr)
    run = p.add_run(text)
    set_run_font(run, size=10.5, color=(0x55, 0x55, 0x55))
    run.font.italic = True
    return p


def generate_report():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    for _ in range(4):
        doc.add_paragraph()
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('Walmart 跨時代財報分析報告')
    set_run_font(run, size=28, bold=True, color=(0x1F, 0x4E, 0x79))
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_before = Pt(12)
    run = subtitle.add_run('以孫子兵法「道・天・地・將・法」五構面\n解讀三十年企業經營智謀')
    set_run_font(run, size=16, color=(0x2E, 0x75, 0xB6))
    doc.add_page_break()

    add_heading_styled(doc, '壹、報告摘要', level=1)
    add_body_text(doc,
        '本報告以 Walmart 兩個歷史時期的合併財務報表為分析基礎——1994 財年代表「實體擴張黃金時代」，'
        '2024 財年代表「數位轉型與全球化時代」——跨越三十年的經營軌跡，從中解讀 Walmart 如何'
        '從一家美國中西部折扣零售商，成長為年營收逾 6,480 億美元的全球零售巨擘。', indent=True)

    add_heading_styled(doc, '【關鍵發現一覽】', level=3)
    add_styled_table(doc, ['指標', '30 年變化'], [
        ['營收規模', '從 679.9 億美元成長至 6,481.3 億美元，30 年間成長約 9.5 倍'],
        ['淨利潤', '從 23.3 億美元成長至 155.1 億美元（歸屬 Walmart），成長約 6.7 倍'],
        ['總資產', '從 264.4 億美元膨脹至 2,524.0 億美元，成長約 9.5 倍'],
        ['營運現金流', '從 22.0 億美元躍升至 357.3 億美元，成長約 16.2 倍'],
        ['流動比率', '從 1.64 下降至 0.83，反映從保守財務走向激進營運資本管理'],
        ['股東權益報酬率', '從 21.7% 下降至 18.5%，規模效應的邊際遞減'],
    ], header_color='1F4E79')
    doc.add_page_break()

    add_heading_styled(doc, '貳、分析框架：孫子兵法五構面', level=1)
    add_quote_box(doc, '「孫子曰：兵者，國之大事，死生之地，存亡之道，不可不察也。'
                       '故經之以五事，校之以計，而索其情：一曰道，二曰天，三曰地，四曰將，五曰法。」'
                       '\n── 《孫子兵法・始計篇》')
    add_styled_table(doc,
        ['五事', '企業構面', '分析內涵', '對應財報指標'],
        [
            ['道（The Way）', '企業使命與經營哲學', '核心價值主張、使命願景', '營收成長、毛利率趨勢、股利政策'],
            ['天（Heaven）', '時勢判斷與外部環境', '總體經濟、科技變革、政策法規', '營收結構變化、新業務佔比'],
            ['地（Earth）', '市場地形與競爭態勢', '市場定位、資產配置、競爭壁壘', '資產結構、資本支出、商譽'],
            ['將（General）', '領導決策與資源配置', '資本配置智慧、併購策略', '股東權益變動、庫藏股、ROE'],
            ['法（Method）', '制度規範與執行效率', '營運效率、現金管理', '現金流量、營運資本、費用率'],
        ], header_color='2E75B6')
    doc.add_page_break()

    add_heading_styled(doc, '參、五構面財務數據對照', level=1)
    add_heading_styled(doc, '【道】營收與毛利率', level=2)
    add_styled_table(doc,
        ['項目', 'FY1994', 'FY1993', 'FY1992', 'FY2024', 'FY2023', 'FY2022'],
        [
            ['淨銷售額（百萬）', '$67,345', '$55,484', '$43,887', '$642,637', '$605,881', '$567,762'],
            ['毛利率', '20.6%', '20.4%', '20.7%', '23.7%', '23.5%', '24.4%'],
            ['淨利率', '3.4%', '3.6%', '3.6%', '2.4%', '1.9%', '2.4%'],
            ['ROE', '21.7%', '22.8%', '—', '18.5%', '15.2%', '—'],
        ], header_color='1F4E79')
    doc.add_paragraph()
    add_heading_styled(doc, '【法】現金流量對照', level=2)
    add_styled_table(doc,
        ['項目', 'FY1994', 'FY1993', 'FY1992', 'FY2024', 'FY2023', 'FY2022'],
        [
            ['營業現金流（百萬）', '$2,195', '$1,278', '$1,357', '$35,726', '$28,841', '$24,181'],
            ['投資現金流', '($4,486)', '($3,506)', '($2,150)', '($21,287)', '($17,722)', '($6,015)'],
            ['自由現金流', '($1,449)', '($2,228)', '($793)', '$15,120', '$11,984', '$11,075'],
        ], header_color='37845D')
    doc.add_page_break()

    add_heading_styled(doc, '肆、五構面綜合評估矩陣', level=1)
    add_styled_table(doc,
        ['構面', 'FY1994 評分', 'FY1994 說明', 'FY2024 評分', 'FY2024 說明'],
        [
            ['道（使命）', '★★★★★', '天天低價 + 快速擴張，使命清晰', '★★★★☆', '價值主張延伸為生態系'],
            ['天（時勢）', '★★★★★', '完美把握實體零售黃金期', '★★★★☆', '成功適應電商衝擊'],
            ['地（地形）', '★★★★☆', '美國本土快速佈點', '★★★★★', '全球佈局完成，實體+數位雙護城河'],
            ['將（領導）', '★★★★★', 'Sam Walton精神傳承，全投入擴張', '★★★★☆', '庫藏股策略靈活'],
            ['法（制度）', '★★★★☆', '衛星通訊領先業界，但現金流為負', '★★★★★', '現金流強勁，供應鏈效率世界級'],
        ], header_color='1F4E79')
    doc.add_page_break()

    add_heading_styled(doc, '伍、結論：三十年的王道啟示', level=1)
    add_quote_box(doc, '「知彼知己，百戰不殆；知天知地，勝乃不窮。」\n── 《孫子兵法》')
    for title_text, body_text in [
        ('啟示一：「道」的恆久性', 'Walmart三十年最成功之處，在於「天天低價」這一核心使命的恆久堅持，但不斷重新定義「低價」的內涵。'),
        ('啟示二：「天」的不可逆性', '從實體零售到全通路經營，Walmart是最堅定的轉型者之一。願意承擔短期成本壓力來佈局未來的企業，才能在下一個時代勝出。'),
        ('啟示三：「地」的累積性', '從264億到2524億的資產規模，Walmart花三十年建立的實體網絡是任何競爭者都無法在短期內複製的。'),
        ('啟示四：「將」的傳承性', '從Sam Walton到職業經理人，Walmart成功實現了領導力的制度化傳承。'),
        ('啟示五：「法」的進化性', '自由現金流從負數轉為正151億的逆轉，是三十年制度建設與效率提升的集大成。'),
    ]:
        add_heading_styled(doc, title_text, level=2)
        add_body_text(doc, body_text, indent=True)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run('【免責聲明】本報告僅供學術研究與教學討論之用，不構成任何投資建議。')
    set_run_font(run, size=9, color=(0x99, 0x99, 0x99))

    output_path = os.path.join(r'c:\Python', 'Walmart_跨時代財報分析報告_王道智謀五構面.docx')
    doc.save(output_path)
    print(f'報告已成功生成：{output_path}')
    return output_path


if __name__ == '__main__':
    generate_report()
